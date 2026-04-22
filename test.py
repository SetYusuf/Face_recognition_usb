import cv2
import os
import pickle
from openpyxl import Workbook, load_workbook
from datetime import datetime

# ---------- CONFIG ----------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

CASCADE_PATH = os.path.join(BASE_DIR, "cascades", "data", "haarcascade_frontalface_alt2.xml")
RECOGNIZER_PATH = os.path.join(BASE_DIR, "trainer.yml")
LABELS_PATH = os.path.join(BASE_DIR, "labels.pickle")

# ---------- EXCEL ----------
excel_file = os.path.join(BASE_DIR, "attendance_clean.xlsx")

def ensure_excel():
    if not os.path.exists(excel_file):
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"
        ws.append(["Name", "Date", "Time"] + [f"Week{i}" for i in range(1, 11)])
        wb.save(excel_file)
        print("Excel created")

ensure_excel()

# ---------- LOAD MODELS ----------
face_cascade = cv2.CascadeClassifier(CASCADE_PATH)
recognizer = cv2.face.LBPHFaceRecognizer_create()
recognizer.read(RECOGNIZER_PATH)

with open(LABELS_PATH, 'rb') as f:
    labels = {v: k for k, v in pickle.load(f).items()}

# ---------- ATTENDANCE ----------
def mark_attendance(name):
    ensure_excel()

    wb = load_workbook(excel_file)
    ws = wb.active

    now = datetime.now()
    date = now.strftime("%Y-%m-%d")
    time = now.strftime("%H:%M:%S")

    # find student row
    row_found = None
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == name:
            row_found = row
            break

    if row_found is None:
        row_found = ws.max_row + 1
        ws.cell(row=row_found, column=1).value = name

    ws.cell(row=row_found, column=2).value = date
    ws.cell(row=row_found, column=3).value = time

    # fill next week slot
    for col in range(4, 14):
        if ws.cell(row=row_found, column=col).value is None:
            ws.cell(row=row_found, column=col).value = 1
            wb.save(excel_file)

            print(f"{name} marked Week {col-3} = 1")
            return

    print(f"{name} already completed 10 weeks")

# ---------- LOOP INPUT ----------
while True:
    IMAGE_PATH = input("\nEnter image path (or q to quit): ").strip()

    if IMAGE_PATH.lower() == "q":
        break

    if not os.path.exists(IMAGE_PATH):
        print("Image not found")
        continue

    img = cv2.imread(IMAGE_PATH)
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)

    faces = face_cascade.detectMultiScale(gray, 1.5, 5)

    if len(faces) == 0:
        print("No face detected")
        continue

    for (x, y, w, h) in faces:
        roi_gray = gray[y:y+h, x:x+w]

        id_, conf = recognizer.predict(roi_gray)

        if 4 <= conf <= 85:
            name = labels[id_]
        else:
            name = "Unknown"

        mark_attendance(name)

        cv2.rectangle(img, (x, y), (x+w, y+h), (255, 0, 0), 2)
        cv2.putText(img, name, (x, y-10),
                    cv2.FONT_HERSHEY_SIMPLEX, 1,
                    (255, 255, 255), 2)

    cv2.imshow("Result", img)
    cv2.waitKey(0)
    cv2.destroyAllWindows()